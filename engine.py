# engine.py - 對帳與報表核心引擎
import sqlite3
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import config

def run_sql_reco(internal_path, broker_path):
    """執行 SQL 全外連接對帳並回傳 Break DataFrame"""
    conn = sqlite3.connect(':memory:')
    
    df_int = pd.read_excel(internal_path)
    df_ext = pd.read_csv(broker_path)

    df_int['Trade_ID'] = df_int['Trade_ID'].astype(str).str.strip().str.upper()
    df_ext['Trade_ID'] = df_ext['Trade_ID'].astype(str).str.strip().str.upper()

    df_int.to_sql('Internal_Trades', conn, index=False, if_exists='replace')
    df_ext.to_sql('External_Trades', conn, index=False, if_exists='replace')

    sql_query = f"""
    SELECT 
        COALESCE(i.Trade_ID, e.Trade_ID) AS Trade_ID,
        i.Product,
        COALESCE(i.Internal_Notional, 0) AS Internal_Notional,
        COALESCE(e.External_Notional, 0) AS External_Notional,
        (COALESCE(i.Internal_Notional, 0) - COALESCE(e.External_Notional, 0)) AS Difference,
        CASE 
            WHEN i.Trade_ID IS NULL THEN 'MISSING_IN_HOUSE'
            WHEN e.Trade_ID IS NULL THEN 'MISSING_ON_BROKER'
            WHEN ABS(COALESCE(i.Internal_Notional, 0) - COALESCE(e.External_Notional, 0)) > {config.TOLERANCE_THRESHOLD} THEN 'NOTIONAL_MISMATCH'
            ELSE 'MATCHED'
        END AS Break_Type
    FROM Internal_Trades i
    FULL OUTER JOIN External_Trades e ON i.Trade_ID = e.Trade_ID
    WHERE Break_Type <> 'MATCHED';
    """

    breaks = pd.read_sql_query(sql_query, conn)
    conn.close()
    return breaks

def apply_excel_styles(output_file):
    """套用 openpyxl 視覺化樣式"""
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active

    fill_red = PatternFill(start_color=config.COLOR_RED, fill_type='solid')
    fill_yellow = PatternFill(start_color=config.COLOR_YELLOW, fill_type='solid')
    fill_orange = PatternFill(start_color=config.COLOR_ORANGE, fill_type='solid')
    fill_header = PatternFill(start_color=config.COLOR_HEADER, fill_type='solid')

    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    font_body = Font(name='Calibri', size=11)
    thin_border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                         top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal='center', vertical='center')

    break_type_col_idx = [cell.value for cell in ws[1]].index('Break_Type') + 1

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        break_type = row[break_type_col_idx - 1].value
        row_fill = fill_red if break_type == 'NOTIONAL_MISMATCH' else (fill_yellow if break_type == 'MISSING_ON_BROKER' else fill_orange)
        for cell in row:
            cell.font = font_body
            cell.border = thin_border
            cell.fill = row_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(output_file)